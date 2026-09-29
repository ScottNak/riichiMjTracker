"""One-time import of past games from MJResults.xlsx into tools/history.json.

Run from the repo folder:  python tools/import_history.py
Then:                      node tools/finish_import.mjs   (checks every game with the app's own code and writes data.json)

Each game tab is read as rounds of point changes. The first row of a round is the payment for the hand;
the rows after it are riichi sticks (-1000 deposits, and the sticks the winner collects).
From those, each round gets its outcome (ron, tsumo, draw, abortive, chombo), who won or dealt in, who was
tenpai, and who declared riichi. The recorded point changes are kept exactly, and the round is marked
imported so the app never recalculates it.

FIXES lists every correction made to the spreadsheet's data, approved by Scott on 2026-09-28.
"""
import json, re, sys
import openpyxl

SOURCE = 'MJResults.xlsx'
OUT = 'tools/history.json'
SKIP = {'PlayerSummary', '(4M-ES) PlayerSummary', 'GameSummary', '4MTemplate', '3MTemplate'}

# Tabs whose date cell is missing or wrong.
DATES = {
    'Aug13': '2024-08-13',  # date cell said 2024-05-07 (May7's date)
    'Jun0325': '2025-06-03', 'Jan1326': '2026-01-13', 'Jan2726': '2026-01-27',
    'Feb1526': '2026-02-15', 'Feb1726': '2026-02-17', 'Mar2426': '2026-03-24',
    'April2126': '2026-04-21', 'April2126 3M': '2026-04-21', 'May19': '2026-05-19',
    'jul1426': '2026-07-14',
    'Jul2725': '2026-07-27',  # tab name typo: 25 for 26
    'sep22 26': '2026-09-22',
}

# East-only games; every other game is East-South.
EAST_ONLY = {'Jan30', 'Feb6', 'Feb20', 'Apr16-1', 'Apr16-2', 'May7', 'May21 (3M)', 'Sep24', 'Oct15',
             'Oct22 (3M)', 'Oct29', 'Nov5', 'Jan2125 (3M)', 'Feb425', 'Feb1125', 'Sep2(3M-E)', 'Sep2(E)'}

NAMES = {'Anvith/Daryl': 'Daryl'}  # Anvith played the first 2 rounds; the game counts for Daryl.

# Corrections to point rows: (tab, row) -> the row's corrected values, in the sheet's column order.
FIXES = {
    ('April2126 3M', 10): [-1000, -1000, 2000],  # E3: Mario won but didn't collect Scott's and Matt's sticks
    ('Jan1326', 27): [-2000, -4000, 6000],       # S2: Scott's mangan tsumo; Ben was dealer and pays 4,000, not Daryl
    ('Apr2925', 30): [5200, -5200, 0, 0],        # S2 (sheet said S2+1, really 0 honba): Scott's ron is 5,200, not 5,500
    ('Feb6', 18): [0, -8000, 8000, 0],           # E4+1: Scott's 7,700 ron was missing the 300 honba
    ('Sep24', 15): [0, -12000, 0, 12000],        # E4: Scott was dealer; 8,000 is impossible, the hand was 12,000
    ('Sep2(E)', 12): [0, 0, -7700, 7700],        # E4: Mahima was dealer; 5,200 becomes the dealer's 7,700
    ('Feb1726', 27): [0, -12000, 12000, 0],      # S3: Scott was dealer; 8,000 becomes the dealer's 12,000
}
# Rows added where a round was missing its stick collection: (tab, row of the round's first line) -> row to add.
ADDED = {
    ('May19', 15): [0, 0, 1000, 0],              # E3+2: Scott's tsumo didn't collect Mario's stick from the draw before
}
# Tabs whose starting dealer isn't the first column: tab -> names from East. Seat order stays as in the columns.
SEATING = {
    'Aug27 (3M)': ['Daryl', 'Kai', 'Scott'],     # the starting dealer wasn't marked; Daryl as East fits every round and the round labels
}
# The dealer +3,000 with everyone else paying 1,000 is either a draw with only the dealer tenpai,
# or a dealer tsumo for 1,000 all: (tab, row) -> 'draw' or 'tsumo'.
DEALER_3000 = {
    ('Feb425', 9): 'tsumo',                      # E3: Scott's dealer tsumo
    ('Feb1726', 3): 'draw',                      # E1: only Rishab tenpai
}
# Rounds with no payments, which can't be read from the numbers: (tab, row) -> outcome.
ZERO_ROUNDS = {
    ('Apr2925', 9): ('abortive', 'kyuushu'),     # E2: kyuushu kyuuhai
    ('Oct22 (3M)', 6): ('draw', None),           # E1+1: everyone tenpai
    ('jul1426', 15): ('draw', None),             # E2: everyone tenpai
}


def num(v):
    return int(round(v)) if isinstance(v, (int, float)) else 0


def read_tab(ws, name):
    players = 4 if isinstance(ws['E2'].value, (int, float)) else 3
    cols = 'BCDE'[:players]
    names = [NAMES.get(str(ws[f'{c}1'].value).strip(), str(ws[f'{c}1'].value).strip()) for c in cols]
    date = DATES.get(name) or ws['A1'].value.strftime('%Y-%m-%d')
    rounds, cur = [], None
    r = 3
    while not str(ws[f'A{r}'].value or '').strip().lower().startswith('final'):
        label = ws[f'A{r}'].value
        if isinstance(ws[f'B{r}'].value, str) and ws[f'B{r}'].value.startswith('=SUM('):
            cur = None  # subtotal row ends a round
        else:
            if label is not None and str(label).strip():
                cur = {'row': r, 'label': str(label).split('\n')[0].strip(), 'lines': []}
                rounds.append(cur)
            values = [ws[f'{c}{r}'].value for c in cols]
            if FIXES.get((name, r)):
                values = FIXES[(name, r)]
            if any(v not in (None, '') for v in values):
                if cur is None:
                    sys.exit(f'{name} row {r}: points with no round label')
                cur['lines'].append([num(v) for v in values])
            if (name, r) in ADDED:
                cur['lines'].append(ADDED[(name, r)])
        r += 1
    rounds = [rd for rd in rounds if rd['lines'] or (name, rd['row']) in ZERO_ROUNDS]
    return {'tab': name, 'date': date, 'names': names, 'players': players, 'rounds': rounds}


# ---- Checking win amounts (used only to pick the seating; finish_import.mjs checks again with the app's code) ----

def ceil100(x):
    return -(-x // 100) * 100


def base_points(han, fu, kiriage):
    if han >= 13: return 8000
    if han >= 11: return 6000
    if han >= 8: return 4000
    if han >= 6: return 3000
    if han >= 5: return 2000
    b = min(fu * 2 ** (han + 2), 2000)
    return 2000 if kiriage and b == 1920 else b


BASES = sorted({base_points(h, f, k) for k in (True, False) for h in range(1, 14)
                for f in [20, 25] + list(range(30, 120, 10))} | {8000 * m for m in range(2, 7)})


def win_fits(kind, pay, winner, dealer, honba, players):
    payers = [s for s in range(players) if s != winner]
    for b in BASES:
        if kind == 'ron':
            if pay[winner] == ceil100((6 if winner == dealer else 4) * b) + 300 * honba:
                return True
        else:
            for each in (100, 150) if players == 3 else (100,):
                if all(-pay[s] == ceil100((2 if winner == dealer or s == dealer else 1) * b) + each * honba for s in payers):
                    return True
    return False


def draw_tenpai(pay, players):
    """The tenpai seats if the payments are a draw's tenpai payments, else None."""
    gainers = [s for s in range(players) if pay[s] > 0]
    losers = [s for s in range(players) if pay[s] < 0]
    if not gainers or not losers or len(gainers) + len(losers) != players:
        return None
    for total in (3000, 2000):
        if all(pay[s] == total // len(gainers) for s in gainers) and all(pay[s] == -total // len(losers) for s in losers):
            return [s in gainers for s in range(players)]
    return None


def read_round(tab, rd, seat_of, players, state):
    """The imported round in seat terms, or an error string."""
    lines = [[0] * players for _ in rd['lines']]
    for i, line in enumerate(rd['lines']):
        for col, v in enumerate(line):
            lines[i][seat_of[col]] = v
    if lines and all(v in (0, -1000) for v in lines[0]):
        pay, sticks = [0] * players, lines
    else:
        pay, sticks = (lines[0] if lines else [0] * players), lines[1:]
    deltas = [sum(line[s] for line in lines) for s in range(players)]
    riichi = [s for s in range(players) if any(line[s] == -1000 for line in sticks)]
    collectors = [s for s in range(players) if any(line[s] > 0 for line in sticks)]
    collected = sum(v for line in sticks for v in line if v > 0)
    if any(v not in (0, -1000) and (v < 0 or v % 1000) for line in sticks for v in line):
        return f"r{rd['row']}: unexpected stick row {sticks}"
    if sum(pay) != 0:
        return f"r{rd['row']}: payments don't add up to 0: {pay}"
    dealer = state['hand'] % players
    on_table = (state['sticks'] + len(riichi)) * 1000
    gainers = [s for s in range(players) if pay[s] > 0]
    losers = [s for s in range(players) if pay[s] < 0]
    base = {'riichi': riichi, 'deltas': deltas, 'imported': True}

    if (tab, rd['row']) in ZERO_ROUNDS:
        outcome, kind = ZERO_ROUNDS[(tab, rd['row'])]
        if any(pay) or collectors:
            return f"r{rd['row']}: expected no payments"
        if outcome == 'abortive':
            return {'outcome': 'abortive', 'kind': kind, **base}
        return {'outcome': 'draw', 'tenpai': [True] * players, **base}
    if collectors:
        if len(collectors) > 1 or gainers != collectors:
            return f"r{rd['row']}: stick collector isn't the only winner: {pay} {sticks}"
        if collected != on_table:
            return f"r{rd['row']}: winner collects {collected} but {on_table} is on the table"
        winner = collectors[0]
        kind = 'ron' if len(losers) == 1 else 'tsumo'
    elif not any(pay):
        return f"r{rd['row']}: no payments and no override"
    else:
        tenpai = draw_tenpai(pay, players)
        if tenpai is not None and gainers == [dealer] and on_table == 0 and win_fits('tsumo', pay, dealer, dealer, state['honba'], players):
            answer = DEALER_3000.get((tab, rd['row']))
            if answer is None:
                return f"r{rd['row']}: could be a draw or a dealer tsumo: {pay}; add it to DEALER_3000"
            if answer == 'draw':
                return {'outcome': 'draw', 'tenpai': tenpai, **base}
        elif tenpai is not None:
            return {'outcome': 'draw', 'tenpai': tenpai, **base}
        if len(losers) == 1 and len(gainers) == players - 1:
            if riichi:
                return f"r{rd['row']}: chombo with riichi sticks"
            return {'outcome': 'chombo', 'offender': losers[0], 'deltas': deltas, 'imported': True}
        if len(gainers) != 1 or losers == []:
            return f"r{rd['row']}: can't read payments {pay}"
        if on_table:
            return f"r{rd['row']}: win with {on_table} on the table not collected"
        winner = gainers[0]
        kind = 'ron' if len(losers) == 1 else 'tsumo'
    if not win_fits(kind, pay, winner, dealer, state['honba'], players):
        return f"r{rd['row']}: {kind} by seat {winner} at {state['honba']} honba fits no score: {pay}"
    rnd = {'outcome': kind, 'winner': winner, **base}
    if kind == 'ron':
        rnd = {'outcome': 'ron', 'winner': winner, 'loser': losers[0], **base}
    return rnd


def advance(state, rnd, players):
    dealer = state['hand'] % players
    s = dict(state)
    if rnd['outcome'] in ('ron', 'tsumo'):
        s['sticks'] = 0
        if rnd['winner'] == dealer: s['honba'] += 1
        else: s['hand'] += 1; s['honba'] = 0
    elif rnd['outcome'] in ('draw', 'abortive'):
        s['sticks'] += len(rnd['riichi'])
        s['honba'] += 1
        if rnd['outcome'] == 'draw' and not rnd['tenpai'][dealer]: s['hand'] += 1
    return s


def read_game(tab):
    """The seating (order[seat] = column) and the rounds in seat terms. Stops at the first round that doesn't read."""
    players = tab['players']
    seating = SEATING.get(tab['tab'])
    order = [tab['names'].index(name) for name in seating] if seating else list(range(players))
    seat_of = {col: seat for seat, col in enumerate(order)}
    state = {'hand': 0, 'honba': 0, 'sticks': 0}
    rounds = []
    for rd in tab['rounds']:
        rnd = read_round(tab['tab'], rd, seat_of, players, state)
        if isinstance(rnd, str):
            sys.exit(f"{tab['tab']} {rnd}")
        rounds.append(rnd)
        state = advance(state, rnd, players)
    return order, rounds


def main():
    wb = openpyxl.load_workbook(SOURCE)
    games = []
    for index, name in enumerate(wb.sheetnames):
        if name in SKIP:
            continue
        tab = read_tab(wb[name], name)
        order, rounds = read_game(tab)
        players = [tab['names'][col] for col in order]
        games.append({
            'id': 'import-' + re.sub(r'[^a-z0-9]+', '-', name.lower()).strip('-'),
            'date': tab['date'],
            'createdAt': f"{tab['date']}T00:00:{index:02d}Z",
            'notes': f'Imported from MJResults.xlsx, tab "{name}".',
            'length': 'east' if name in EAST_ONLY else 'south',
            'players': players,
            'rounds': rounds,
        })
        seating = '' if list(order) == list(range(tab['players'])) else f"  (seated {', '.join(players)})"
        print(f"{name:14} {tab['date']}  {len(rounds):2} rounds{seating}")
    with open(OUT, 'w', encoding='utf-8') as f:
        json.dump(games, f, indent=1)
    print(f'{len(games)} games written to {OUT}')


main()
